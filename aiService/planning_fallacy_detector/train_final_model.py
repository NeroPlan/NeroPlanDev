"""
최종 모델 학습 + 저장 (시연용)
- 페르소나 20개
- 전체 데이터로 최종 학습(교차검증 아님) 후 joblib으로 저장 -> 데모 스크립트에서 바로 로드해서 씀
- 모든 성능 지표는 점추정치 + 페르소나 클러스터 붓스트랩 95% CI를 함께 보고
- 확인용 데이터는 CSV 대신 엑셀(.xlsx)로 저장: data 시트(원본 행) + 컬럼설명 시트(컬럼 뜻 + 산출식)를 한 번에 생성
"""
import joblib
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font
from bias_data_pipeline import make_random_personas, generate_raw_events, add_features, build_pipeline, bootstrap_ci
from sklearn.metrics import roc_auc_score, average_precision_score, precision_recall_curve
from sklearn.model_selection import GroupKFold, cross_val_predict

N_PERSONAS = 20
SEED = 42
N_BOOTSTRAP = 2000


def save_data_workbook(data: pd.DataFrame, path: str = "synthetic_training_data.xlsx"):
    """확인용 엑셀 저장: data 시트(원본 행) + 컬럼설명 시트(컬럼 뜻 + 산출식)."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        data.to_excel(writer, sheet_name="data", index=False)

        columns_info = [
            ("persona_id", "어떤 가상 사용자(페르소나)의 기록인지. persona_0이면 0번째로 생성된 페르소나."),
            ("task_index", "그 페르소나 안에서 몇 번째 과제인지(0부터 시작, 시간 순서)"),
            ("category", "과제 카테고리(운동/공부/업무/개인일정)"),
            ("time_slot", "과제 시간대(아침/오후/저녁/밤)"),
            ("estimated_minutes", "예상 소요시간(30/60/90/120분 중 하나)"),
            ("outcome", "실제 결과(완료/미룸/실패)"),
            ("y", "모델 학습용 타겟값. outcome이 완료면 0, 미룸이나 실패면 1(둘 다 \"안 지켜짐\"으로 묶어서 1)"),
            ("category_relative_rate", "이 과제 시점까지, 이 카테고리 완료율이 본인 전체 완료율보다 얼마나 높거나 낮은지(최근 15개 이동창 기준)"),
            ("recent_failure_streak", "이 과제 시점까지, 같은 카테고리에서 직전까지 연속으로 몇 번 안 지켰는지"),
        ]
        df_columns = pd.DataFrame(columns_info, columns=["컬럼", "뜻"])
        df_columns.to_excel(writer, sheet_name="컬럼설명", index=False, startrow=0)

        formula_rows = [
            "",
            "category_relative_rate 산출식",
            "category_relative_rate = cat_rate − overall_rate",
            "  cat_rate = mean(이 시점 직전까지, 같은 카테고리에서 최근 15건 이내의 완료여부(완료=1, 그외=0))",
            "  overall_rate = mean(이 시점 직전까지, 카테고리 무관 전체 최근 15건 이내의 완료여부(완료=1, 그외=0))",
            "  단, cat_rate 또는 overall_rate를 계산할 이력이 하나도 없으면(콜드스타트) → category_relative_rate = 0",
            "",
            "recent_failure_streak 산출식",
            "recent_failure_streak = 직전 같은 카테고리 과제부터 거슬러 올라가며 연속으로 '완료가 아님'이 이어진 횟수",
            "  완료(outcome=완료)가 한 번이라도 나오면 그 시점에서 0으로 리셋되고, 그 다음부터 다시 카운트",
            "  이력이 없으면(콜드스타트) → recent_failure_streak = 0",
        ]
        ws = writer.sheets["컬럼설명"]
        start = len(df_columns) + 3
        for i, text in enumerate(formula_rows):
            if text:
                ws.cell(row=start + i, column=1, value=text)

        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 110
        for row in ws.iter_rows(min_row=1, max_row=len(df_columns) + 1, min_col=2, max_col=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in ws[1]:
            cell.font = Font(bold=True)

    print(f"합성데이터 저장 완료: {path} ({len(data)}행, 시트: data, 컬럼설명)")


personas = make_random_personas(N_PERSONAS, seed=SEED)
raw = generate_raw_events(personas, seed=SEED + 1000)
data = add_features(raw)

# 확인용: 학습에 쓴 합성데이터 원본을 엑셀로 저장 (joblib엔 안 들어있는 원본 행 확인용)
save_data_workbook(data)

y = data["y"].values
groups = data["persona_id"].values
X = data[["estimated_minutes", "category_relative_rate", "recent_failure_streak", "category", "time_slot"]]

# 교차검증으로 임계값용 out-of-fold 확률 생성 (성능 평가용)
# 4묶음 학습하고 안 본 1묶음으로 예측 - 5번 반복
gkf = GroupKFold(n_splits=5)
probs_oof = cross_val_predict(build_pipeline(), X, y, cv=gkf.split(X, y, groups), method="predict_proba")[:, 1]

precisions, recalls, thresholds = precision_recall_curve(y, probs_oof)
valid_idx = np.where(precisions[:-1] >= 0.8)[0]
threshold = thresholds[valid_idx[np.argmin(thresholds[valid_idx])]] if len(valid_idx) else 0.5

# 점추정치 + 페르소나 클러스터 붓스트랩 95% CI (EBSLN 논문 Table 5 방식)
auc, auc_boot_mean, auc_ci_low, auc_ci_high = bootstrap_ci(groups, y, probs_oof, roc_auc_score, n_boot=N_BOOTSTRAP, seed=1)
ap, ap_boot_mean, ap_ci_low, ap_ci_high = bootstrap_ci(groups, y, probs_oof, average_precision_score, n_boot=N_BOOTSTRAP, seed=2)

print(f"[페르소나 {N_PERSONAS}개 기준] OOF 성능 (점추정치, 95% CI)")
print(f"  AUC = {auc:.3f}  (95% CI {auc_ci_low:.3f}-{auc_ci_high:.3f})")
print(f"  AP  = {ap:.3f}  (95% CI {ap_ci_low:.3f}-{ap_ci_high:.3f})")
print(f"  채택임계값 = {threshold:.3f} (precision-recall curve, precision>=0.8 최저점)")

# 최종 배포용: 전체 데이터로 다시 학습(교차검증은 검증용, 최종 모델은 전체 데이터 활용)
final_model = build_pipeline()
final_model.fit(X, y)

joblib.dump({
    "model": final_model,
    "threshold": float(threshold),
    "auc": float(auc), "auc_ci": [float(auc_ci_low), float(auc_ci_high)],
    "ap": float(ap), "ap_ci": [float(ap_ci_low), float(ap_ci_high)],
    "n_personas": N_PERSONAS,
}, "final_bias_model.joblib")
print("저장 완료: final_bias_model.joblib")